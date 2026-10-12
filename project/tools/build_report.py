# -*- coding: utf-8 -*-
import os
import re
import shutil
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = r"C:\Users\马语泽\Desktop\案例网页分析.docx"
TMP = os.path.join(os.path.dirname(__file__), "_report_build.docx")

EAST = "微软雅黑"
LATIN = "Calibri"
BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x38, 0x64)
GRAY = RGBColor(0x59, 0x59, 0x59)
HEADER_GRAY = RGBColor(0x80, 0x80, 0x80)


def set_font(style, size, color=None, bold=None, east=EAST, latin=LATIN):
    style.font.name = latin
    style.font.size = Pt(size)
    if color is not None:
        style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east)


def add_page_field(paragraph, size=9, color=HEADER_GRAY):
    run = paragraph.add_run()
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = LATIN
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), EAST)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._element.append(fld_begin)
    run._element.append(instr)
    run._element.append(fld_end)


doc = Document()

sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(1)
sec.bottom_margin = Inches(1)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)
sec.header_distance = Inches(0.492)
sec.footer_distance = Inches(0.492)

normal = doc.styles["Normal"]
set_font(normal, 11)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

title_style = doc.styles["Title"]
set_font(title_style, 22, DARK_BLUE, True)
title_style.paragraph_format.space_before = Pt(0)
title_style.paragraph_format.space_after = Pt(6)

subtitle_style = doc.styles["Subtitle"]
set_font(subtitle_style, 12, GRAY, False)
subtitle_style.paragraph_format.space_before = Pt(0)
subtitle_style.paragraph_format.space_after = Pt(12)

h1_style = doc.styles["Heading 1"]
set_font(h1_style, 16, BLUE, True)
h1_style.paragraph_format.space_before = Pt(16)
h1_style.paragraph_format.space_after = Pt(8)
h1_style.paragraph_format.keep_with_next = True

h2_style = doc.styles["Heading 2"]
set_font(h2_style, 13, BLUE, True)
h2_style.paragraph_format.space_before = Pt(12)
h2_style.paragraph_format.space_after = Pt(6)
h2_style.paragraph_format.keep_with_next = True

bullet_style = doc.styles["List Bullet"]
set_font(bullet_style, 11)
bullet_style.paragraph_format.space_before = Pt(0)
bullet_style.paragraph_format.space_after = Pt(8)
bullet_style.paragraph_format.line_spacing = 1.167
bullet_style.paragraph_format.left_indent = Inches(0.5)
bullet_style.paragraph_format.first_line_indent = Inches(-0.25)

header = sec.header
hp = header.paragraphs[0]
hp.text = "案例网页分析报告"
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
for run in hp.runs:
    run.font.size = Pt(9)
    run.font.color.rgb = HEADER_GRAY
    run.font.name = LATIN
p_pr = hp._p.get_or_add_pPr()
p_bdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single")
bottom.set(qn("w:sz"), "4")
bottom.set(qn("w:space"), "4")
bottom.set(qn("w:color"), "BFBFBF")
p_bdr.append(bottom)
p_pr.append(p_bdr)

footer = sec.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
add_page_field(fp)


def h1(text):
    return doc.add_paragraph(text, style="Heading 1")


def h2(text):
    return doc.add_paragraph(text, style="Heading 2")


def body(text):
    return doc.add_paragraph(text)


def bullet(text):
    return doc.add_paragraph(text, style="List Bullet")


doc.core_properties.title = "案例网页分析报告"
doc.core_properties.subject = "滚动驱动叙事网页技术路线分析"
doc.core_properties.author = "焗瓷记项目组"

doc.add_paragraph("案例网页分析报告", style="Title")
doc.add_paragraph("滚动驱动叙事网页技术路线分析", style="Subtitle")

h1("一、核心内容载体：基于静态图片序列")
body("该网页的核心交互体验（滚动控制叙事进程）并非基于传统视频文件，而是基于预先渲染并优化好的静态图片序列。")
bullet("证据：网页源代码中未发现 <video> 标签或对 .mp4、.webm 等视频资源的引用。其滚动驱动的视觉变化，是通过快速切换一系列连贯的、高画质的 JPEG 或 WebP 格式图片来实现的，每一张图片对应动画的一个关键帧。")
bullet("优势：这种方式能确保每一帧的画面质量都绝对清晰、稳定，且不受网络带宽和解码性能影响，可实现“帧-帧”级别的精确控制，完美支持顺滑的滚动驱动和惯性效果。")

h1("二、主要技术路线解析")
body("该网站采用了一套典型且成熟的“滚动驱动动画”技术方案，其核心流程如下：")

h2("前端交互层（JavaScript）")
bullet("监听事件：使用 addEventListener 监听鼠标的 wheel 滚动事件。")
bullet("计算进度：在事件回调中，计算滚动距离（deltaY），并将其映射为动画的整体进度百分比（0% 到 100%）。")
bullet("控制播放：根据进度值，动态更新显示对应的图片帧（通过改变 <img> 标签的 src 属性或 background-image）。")

h2("惯性缓动与状态管理")
bullet("为实现“停止滚动后继续平滑移动”的效果，页面并未直接使用 currentTime，而是运用了缓动函数（easing function）。")
bullet("它维护一个“当前显示帧”和“目标滚动帧”。当滚动停止，系统通过 requestAnimationFrame 启动一个缓动动画循环，让“当前帧”以逐渐减小的速度向“目标帧”靠拢，从而模拟出物理惯性。")

h2("内容组织与性能优化")
bullet("图片预加载：关键图片序列会提前预加载，确保滚动时画面切换流畅无白屏。")
bullet("懒加载：非关键或后续章节的图片可能采用懒加载策略，优化首屏加载速度。")
bullet("CSS 与设计：配合精细的 CSS 过渡和排版，确保图文、文物特写与背景动画的视觉统一。")

h1("三、与你需求的关键匹配点")
body("此技术路线与你之前提出的需求（基于视频 网页端主体视频.webm）在核心实现逻辑上是一致的，主要区别在于内容素材的形态：")

h2("相同点")
body("交互控制的核心逻辑（滚动驱动、惯性缓动、进度映射）完全通用。")

h2("差异与适配")
body("该网页使用的是“图片序列”，而你的素材是“视频文件”。")
bullet("直接方法：你可以将视频按固定帧率（如 25fps 或 30fps）导出为图片序列，然后完全复刻该网站的技术路线，这是实现其“丝滑惯性”效果的最直接方式。")
bullet("替代方法：也可直接控制 <video> 元素的 currentTime 属性，但需自行实现惯性逻辑，且效果可能受视频编码影响。")

h1("四、技术栈推测")
bullet("核心语言：HTML5、CSS3、JavaScript（ES6+）。")
bullet("关键 API：wheel 事件、requestAnimationFrame、Image 对象预加载。")
bullet("构建工具：可能使用了 Webpack 或 Vite 等模块打包工具。")
bullet("框架：从页面结构看，可能基于 React 或 Vue 等现代前端框架开发，但滚动驱动的核心逻辑是框架无关的。")

doc.save(TMP)


def patch_numbering():
    with zipfile.ZipFile(TMP, "r") as zin:
        items = {name: zin.read(name) for name in zin.namelist()}

    numbering = items["word/numbering.xml"].decode("utf-8")
    match = re.search(r'(<w:abstractNum w:abstractNumId="8">.*?</w:abstractNum>)', numbering, re.S)
    if match:
        block = match.group(1)
        patched = block.replace('<w:tab w:val="num" w:pos="360"/>', '<w:tab w:val="num" w:pos="720"/>')
        patched = patched.replace('<w:ind w:left="360" w:hanging="360"/>', '<w:ind w:left="720" w:hanging="360"/>')
        numbering = numbering.replace(block, patched)
    items["word/numbering.xml"] = numbering.encode("utf-8")

    with zipfile.ZipFile(TMP, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, blob in items.items():
            zout.writestr(name, blob)


patch_numbering()

if os.path.exists(OUT):
    os.remove(OUT)
shutil.move(TMP, OUT)
print("written:", OUT)
