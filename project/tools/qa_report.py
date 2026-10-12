# -*- coding: utf-8 -*-
import re
import zipfile

from docx import Document

p = r"C:\Users\马语泽\Desktop\案例网页分析.docx"
z = zipfile.ZipFile(p)
print("entries:", len(z.namelist()))
s = z.read("word/document.xml").decode("utf-8")
print("title:", "案例网页分析报告" in s)
print("listbullet style used:", 'w:val="ListBullet"' in s)
print("listbullet count:", s.count('w:val="ListBullet"'))
n = z.read("word/numbering.xml").decode("utf-8")
m = re.search(r'<w:abstractNum w:abstractNumId="8">.*?</w:abstractNum>', n, re.S)
print("num8 tab720:", 'pos="720"' in m.group(0))
print("num8 left720:", 'left="720"' in m.group(0))
st = z.read("word/styles.xml").decode("utf-8")
print("calibri:", "Calibri" in st)
print("yahei:", "微软雅黑" in st)
print("h1 color:", "2E74B5" in st)
print("bullet after160:", 'w:after="160"' in st)
print("sect margins:", 'w:left="1440"' in s)

doc = Document(p)
styles_used = {}
for para in doc.paragraphs:
    styles_used[para.style.name] = styles_used.get(para.style.name, 0) + 1
print("paragraph count:", len(doc.paragraphs))
print("styles used:", styles_used)
print("first para:", doc.paragraphs[0].text)
print("last para:", doc.paragraphs[-1].text[:30])
print("header text:", doc.sections[0].header.paragraphs[0].text)
print("footer field:", "PAGE" in doc.sections[0].footer.paragraphs[0]._p.xml)
print("title prop:", doc.core_properties.title)
